"""議事録の Markdown / DOCX エクスポート（仕様セクション4のフォーマット）。"""
import io

from ..schemas import MinutesContent


def render_markdown(content: MinutesContent) -> str:
    g = content.kaigi_gaiyou
    lines: list[str] = ["# 議事録", "", "## 会議概要"]
    lines += [
        f"- 会議名：{g.kaigi_mei}",
        f"- 日時：{g.nichiji}",
        f"- 場所：{g.basho}",
        f"- 出席者：{'、'.join(g.shussekisha)}",
        f"- 欠席者：{'、'.join(g.kessekisha)}",
        "- 書記：AI議事録エージェント",
        "",
        "## 議題",
    ]
    lines += [f"{i}. {item}" for i, item in enumerate(content.gidai, 1)] or ["（なし）"]

    lines += ["", "## 議論内容"]
    for topic in content.giron_naiyou:
        lines.append(f"### {topic.gidai}")
        for y in topic.youten:
            speaker = f"（{y.hatsugensha}）" if y.hatsugensha else ""
            lines.append(f"- {y.naiyou}{speaker}")
        lines.append("")

    lines += ["## 決定事項"]
    lines += [f"- {d}" for d in content.kettei_jikou] or ["- （なし）"]

    lines += ["", "## アクションアイテム", "| No. | 内容 | 担当者 | 期限 |", "|-----|------|--------|------|"]
    lines += [
        f"| {a.no} | {a.naiyou} | {a.tantousha} | {a.kigen} |" for a in content.action_items
    ]

    lines += [
        "",
        "## 次回会議",
        f"- 日時：{content.jikai_kaigi.nichiji}",
        "- 議題（予定）：" + "、".join(content.jikai_kaigi.gidai_yotei),
        "",
        "## 保留・継続検討事項",
    ]
    lines += [f"- {h}" for h in content.horyuu_jikou] or ["- （なし）"]
    return "\n".join(lines) + "\n"


def render_docx(content: MinutesContent) -> bytes:
    from docx import Document  # 遅延 import

    doc = Document()
    g = content.kaigi_gaiyou
    doc.add_heading("議事録", level=0)

    doc.add_heading("会議概要", level=1)
    for label, value in [
        ("会議名", g.kaigi_mei),
        ("日時", g.nichiji),
        ("場所", g.basho),
        ("出席者", "、".join(g.shussekisha)),
        ("欠席者", "、".join(g.kessekisha)),
        ("書記", "AI議事録エージェント"),
    ]:
        doc.add_paragraph(f"{label}：{value}", style="List Bullet")

    doc.add_heading("議題", level=1)
    for i, item in enumerate(content.gidai, 1):
        doc.add_paragraph(f"{i}. {item}")

    doc.add_heading("議論内容", level=1)
    for topic in content.giron_naiyou:
        doc.add_heading(topic.gidai, level=2)
        for y in topic.youten:
            speaker = f"（{y.hatsugensha}）" if y.hatsugensha else ""
            doc.add_paragraph(f"{y.naiyou}{speaker}", style="List Bullet")

    doc.add_heading("決定事項", level=1)
    for d in content.kettei_jikou:
        doc.add_paragraph(d, style="List Bullet")

    doc.add_heading("アクションアイテム", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["No.", "内容", "担当者", "期限"]):
        cell.text = text
    for a in content.action_items:
        row = table.add_row().cells
        row[0].text = str(a.no)
        row[1].text = a.naiyou
        row[2].text = a.tantousha
        row[3].text = a.kigen

    doc.add_heading("次回会議", level=1)
    doc.add_paragraph(f"日時：{content.jikai_kaigi.nichiji}", style="List Bullet")
    doc.add_paragraph(
        "議題（予定）：" + "、".join(content.jikai_kaigi.gidai_yotei), style="List Bullet"
    )

    doc.add_heading("保留・継続検討事項", level=1)
    for h in content.horyuu_jikou:
        doc.add_paragraph(h, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
